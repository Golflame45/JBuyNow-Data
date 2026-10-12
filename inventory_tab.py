import io
import os
import re
import pandas as pd
import numpy as np
import plotly.graph_objects as go
import streamlit as st
from datetime import datetime

# ================= 1. Data Preparation & Caching =================
@st.cache_data(ttl=300)
def get_inventory_valuation_data(master_df_summary, max_date_str, active_stock_name=""):
    """
    Builds the complete inventory valuation dataset by joining:
    1. Active Stock File (from session state or local/GDrive)
    2. SKU Master File (Categories, Description, Cost per unit)
    3. Historical Sales Performance (Lifetime Realized Avg Price & 30-Day Velocity)
    """
    base_dir = os.path.dirname(os.path.abspath(__file__)) if '__file__' in globals() else '.'
    
    # 1. Load Active Stock File
    stock_df = None
    if 'raw_stock_df' in st.session_state and not st.session_state['raw_stock_df'].empty:
        stock_df = st.session_state['raw_stock_df'].copy()
    else:
        # Search candidate stock files
        stock_candidates = [
            os.path.join(base_dir, 'Stock as of 09-10.xlsx'),
            os.path.join(base_dir, 'Stock as of 07.10.xlsx'),
            os.path.join(base_dir, 'Stock as of 3.10.xlsx'),
            os.path.join(base_dir, 'Stock as of 25.09.xlsx'),
            r"C:\Users\user\OneDrive\Personal\Consultant\JBuynow\Test Data\Stock as of 09-10.xlsx",
            r"C:\Users\user\OneDrive\Personal\Consultant\JBuynow\Test Data\Stock as of 07.10.xlsx"
        ]
        target_stock_path = next((p for p in stock_candidates if os.path.exists(p)), None)
        if target_stock_path:
            try:
                raw_s = pd.read_excel(target_stock_path)
                sku_c = next((c for c in raw_s.columns if str(c).strip().lower() in ['no.', 'no', 'item no.', 'sku', 'seller sku', 'รหัสสินค้า']), None)
                stk_c = next((c for c in raw_s.columns if any(k in str(c).strip().lower() for k in ['available', 'inventory', 'stock', 'on hand', 'qty', 'สต๊อก', 'พร้อมขาย'])), None)
                if sku_c and stk_c:
                    raw_s = raw_s[~raw_s[sku_c].astype(str).str.contains('ผลรวม|total|grand', case=False, na=False)]
                    raw_s = raw_s.dropna(subset=[sku_c])
                    raw_s['SKU'] = raw_s[sku_c].astype(str).str.strip()
                    raw_s['Stock_Available'] = pd.to_numeric(raw_s[stk_c].astype(str).str.replace(',', '').str.replace('-', '0'), errors='coerce').fillna(0)
                    stock_df = raw_s.groupby('SKU')['Stock_Available'].sum().reset_index()
            except Exception:
                pass

    if stock_df is None or stock_df.empty:
        return pd.DataFrame()

    # Clean Stock DF
    stock_df = stock_df[~stock_df['SKU'].astype(str).str.contains('ผลรวม|total|grand', case=False, na=False)]
    stock_df = stock_df.dropna(subset=['SKU'])
    stock_df['SKU'] = stock_df['SKU'].astype(str).str.strip()
    stock_df['Stock_Available'] = pd.to_numeric(stock_df['Stock_Available'], errors='coerce').fillna(0)
    stock_table = stock_df.groupby('SKU', as_index=False)['Stock_Available'].sum()

    # 2. Load Master SKU File
    sku_master_df = None
    if 'sku_master_df' in st.session_state and isinstance(st.session_state['sku_master_df'], pd.DataFrame) and not st.session_state['sku_master_df'].empty:
        sku_master_df = st.session_state['sku_master_df'].copy()
    else:
        sku_candidates = [
            os.path.join(base_dir, 'Master SKU.xlsx'),
            r"C:\Users\user\OneDrive\Personal\Consultant\JBuynow\Test Data\Master SKU.xlsx"
        ]
        target_sku_path = next((p for p in sku_candidates if os.path.exists(p)), None)
        if target_sku_path:
            try:
                sku_master_df = pd.read_excel(target_sku_path)
            except Exception:
                pass

    if sku_master_df is not None and not sku_master_df.empty:
        sku_col = next((c for c in sku_master_df.columns if str(c).strip().lower() in ['no.', 'no', 'item no.', 'sku', 'seller sku', 'รหัสสินค้า']), None)
        desc_col = next((c for c in sku_master_df.columns if str(c).strip().lower() in ['description', 'ชื่อสินค้า', 'name', 'product']), None)
        cat_col = next((c for c in sku_master_df.columns if str(c).strip().lower() in ['category description', 'cate desc', 'category_desc', 'หมวดหมู่สินค้า', 'หมวดหมู่']), None)
        pcat_col = next((c for c in sku_master_df.columns if str(c).strip().lower() in ['parent category description', 'parent category', 'หมวดหมู่หลัก']), None)
        cost_col = next((c for c in sku_master_df.columns if str(c).strip().lower() in ['bkk jlo', 'cost', 'ต้นทุน', 'ทุน', 'jlo']), None)

        if sku_col:
            sku_master_df['SKU'] = sku_master_df[sku_col].astype(str).str.strip()
            sku_master_df['Description'] = sku_master_df[desc_col].astype(str).str.strip() if desc_col else sku_master_df['SKU']
            sku_master_df['Category_Desc'] = sku_master_df[cat_col].astype(str).str.strip() if cat_col else 'อื่นๆ / ไม่ระบุหมวด'
            sku_master_df['Parent_Category'] = sku_master_df[pcat_col].astype(str).str.strip() if pcat_col else 'อื่นๆ'
            sku_master_df['Cost_Unit'] = pd.to_numeric(sku_master_df[cost_col], errors='coerce').fillna(0.0) if cost_col else 0.0

            # Build prefix fallback dictionaries
            sku_master_df['SKU_Prefix'] = sku_master_df['SKU'].str.split('-').str[0]
            prefix_cat_map = sku_master_df.groupby('SKU_Prefix')['Category_Desc'].first().to_dict()
            prefix_pcat_map = sku_master_df.groupby('SKU_Prefix')['Parent_Category'].first().to_dict()
            prefix_cost_map = sku_master_df.groupby('SKU_Prefix')['Cost_Unit'].mean().to_dict()

            sku_info = sku_master_df[['SKU', 'Description', 'Category_Desc', 'Parent_Category', 'Cost_Unit']].drop_duplicates(subset=['SKU'])
            stock_table = stock_table.merge(sku_info, on='SKU', how='left')

            # Intelligent Prefix Fallback for new variants or unmatched items
            stock_table['SKU_Prefix'] = stock_table['SKU'].str.split('-').str[0]
            stock_table['Category_Desc'] = stock_table['Category_Desc'].fillna(stock_table['SKU_Prefix'].map(prefix_cat_map)).fillna('อื่นๆ / ไม่ระบุหมวด')
            stock_table['Parent_Category'] = stock_table['Parent_Category'].fillna(stock_table['SKU_Prefix'].map(prefix_pcat_map)).fillna('อื่นๆ')
            stock_table['Cost_Unit'] = stock_table['Cost_Unit'].fillna(stock_table['SKU_Prefix'].map(prefix_cost_map)).fillna(0.0)
            stock_table['Description'] = stock_table['Description'].fillna(stock_table['SKU'])
            stock_table.drop(columns=['SKU_Prefix'], inplace=True, errors='ignore')
    else:
        stock_table['Description'] = stock_table['SKU']
        stock_table['Category_Desc'] = 'อื่นๆ / ไม่ระบุหมวด'
        stock_table['Parent_Category'] = 'อื่นๆ'
        stock_table['Cost_Unit'] = 0.0

    # 3. Merge Sales Performance Summary
    if master_df_summary is not None and not master_df_summary.empty:
        stock_table = stock_table.merge(master_df_summary, on='SKU', how='left')
    
    for c in ['Avg_Selling_Price', 'Units_Sold', 'Revenue', 'Units_Sold_30D']:
        if c not in stock_table.columns:
            stock_table[c] = 0.0
        else:
            stock_table[c] = pd.to_numeric(stock_table[c], errors='coerce').fillna(0.0)

    # 4. Values and Profit Margins
    stock_table['Total_Cost_Value'] = stock_table['Stock_Available'] * stock_table['Cost_Unit']
    stock_table['Expected_Sales_Value'] = stock_table['Stock_Available'] * stock_table['Avg_Selling_Price']
    stock_table['Projected_Margin'] = np.maximum(0.0, stock_table['Expected_Sales_Value'] - stock_table['Total_Cost_Value'])

    # 5. Inventory Velocity & Health Status Classification
    stock_table['Daily_Run_Rate'] = stock_table['Units_Sold_30D'] / 30.0
    stock_table['Days_of_Inventory'] = np.where(
        stock_table['Daily_Run_Rate'] > 0,
        stock_table['Stock_Available'] / stock_table['Daily_Run_Rate'],
        999.0
    )

    def classify_stock(row):
        stk = row['Stock_Available']
        days = row['Days_of_Inventory']
        sold_30d = row['Units_Sold_30D']
        if stk <= 0:
            return '⚪ หมดสต็อก (0 ชิ้น)'
        elif sold_30d > 0 and days <= 14:
            return '🔴 สินค้าใกล้หมด (< 14 วัน)'
        elif sold_30d > 0 and days <= 90:
            return '🟢 สต็อกปกติ (14-90 วัน)'
        elif sold_30d > 0 and days > 90:
            return '🟠 สต็อกล้น (> 90 วัน)'
        elif sold_30d == 0 and stk > 0:
            return '🟡 สินค้าค้างสต็อก (ไม่มียอดขาย 30 วัน)'
        else:
            return '🟢 สต็อกปกติ (14-90 วัน)'

    stock_table['Stock_Status'] = stock_table.apply(classify_stock, axis=1)
    
    # Sort deterministically by Expected Sales Value
    stock_table = stock_table.sort_values(['Expected_Sales_Value', 'Stock_Available'], ascending=[False, False]).reset_index(drop=True)
    return stock_table


# ================= 2. Excel Export Helper =================
def export_inventory_to_excel(df_inv):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        cols = [
            'SKU', 'Description', 'Category_Desc', 'Stock_Status',
            'Stock_Available', 'Cost_Unit', 'Total_Cost_Value',
            'Avg_Selling_Price', 'Expected_Sales_Value', 'Units_Sold_30D', 'Days_of_Inventory'
        ]
        actual_cols = [c for c in cols if c in df_inv.columns]
        out_df = df_inv[actual_cols].copy()
        rename_map = {
            'SKU': 'รหัสสินค้า (SKU)',
            'Description': 'ชื่อสินค้า / รายละเอียด',
            'Category_Desc': 'หมวดหมู่สินค้า (Category)',
            'Stock_Status': 'สถานะสุขภาพสต็อก',
            'Stock_Available': 'สต็อกคงเหลือ (ชิ้น)',
            'Cost_Unit': 'ต้นทุน/ชิ้น (฿)',
            'Total_Cost_Value': 'มูลค่าต้นทุนรวม (฿)',
            'Avg_Selling_Price': 'ราคาขายเฉลี่ย (฿)',
            'Expected_Sales_Value': 'มูลค่าขายที่คาดหวัง (฿)',
            'Units_Sold_30D': 'ยอดขาย 30 วันล่าสุด (ชิ้น)',
            'Days_of_Inventory': 'สต็อกพอขาย (วัน)'
        }
        out_df = out_df.rename(columns=rename_map)
        out_df.to_excel(writer, index=False, sheet_name='Stock_Valuation')
        ws = writer.sheets['Stock_Valuation']
        try:
            from openpyxl.styles import Font, PatternFill, Alignment
            header_fill = PatternFill(start_color='1F2937', end_color='1F2937', fill_type='solid')
            header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)
        except Exception:
            pass
    buf.seek(0)
    return buf.getvalue()


# ================= 3. Inventory Dashboard Component =================
@st.fragment
def render_inventory_valuation_tab(base_df, active_stock_name=""):
    """
    Renders the Executive Inventory & Valuation Dashboard inside the Stock Tab.
    """
    if base_df is None or base_df.empty:
        st.info("ไม่มีข้อมูลฐานยอดขายสำหรับคำนวณราคาขายเฉลี่ย")
        return

    # Extract Sales Summaries (Fast Pre-aggregation)
    sales_all = base_df.groupby('SKU').agg({
        'Revenue': 'sum',
        'Units_Sold': 'sum',
        'Orders': 'sum'
    }).reset_index()
    sales_all['Avg_Selling_Price'] = np.where(sales_all['Units_Sold'] > 0, sales_all['Revenue'] / sales_all['Units_Sold'], 0.0)

    # 30-Day Sales Velocity
    max_date = base_df['DateObj'].max() if 'DateObj' in base_df.columns else pd.Timestamp.now()
    last_30_date = max_date - pd.Timedelta(days=30)
    df_recent = base_df[base_df['DateObj'] >= last_30_date] if 'DateObj' in base_df.columns else pd.DataFrame()
    
    if not df_recent.empty:
        recent_sales = df_recent.groupby('SKU')['Units_Sold'].sum().reset_index().rename(columns={'Units_Sold': 'Units_Sold_30D'})
        sales_summary = sales_all.merge(recent_sales, on='SKU', how='left')
    else:
        sales_summary = sales_all.copy()
        sales_summary['Units_Sold_30D'] = 0.0

    sales_summary['Units_Sold_30D'] = sales_summary['Units_Sold_30D'].fillna(0.0)

    # Load Full Inventory Table
    max_d_str = str(max_date)[:10]
    stock_table = get_inventory_valuation_data(sales_summary, max_d_str, active_stock_name)

    if stock_table.empty:
        st.warning("⚠️ ยังไม่พบไฟล์ข้อมูลสต็อกสินค้าในระบบ กรุณาตรวจสอบไฟล์สต็อกใน Google Drive ห้อง 02_Master_Data ครับ")
        return

    # Subtitle Header
    disp_stock_name = active_stock_name if active_stock_name else "Stock as of 09-10.xlsx"
    st.markdown(f"""
        <div style="background-color: #1e293b; padding: 10px 16px; border-radius: 8px; border-left: 4px solid #3b82f6; margin-bottom: 15px; display: flex; justify-content: space-between; align-items: center;">
            <span style="font-size: 15px; font-weight: 600; color: #f8fafc;">📦 ระบบบริหารสต็อกสินค้าและมูลค่าคงคลัง (Inventory Valuation & Health)</span>
            <span style="font-size: 13px; color: #94a3b8; background: #0f172a; padding: 4px 10px; border-radius: 4px; border: 1px solid #334155;">📄 ไฟล์สต็อก: <b>{disp_stock_name}</b></span>
        </div>
    """, unsafe_allow_html=True)

    # ================= Filters Row =================
    col_f1, col_f2, col_f3 = st.columns([1.8, 1.8, 2.4])

    with col_f1:
        all_categories = sorted([str(c) for c in stock_table['Category_Desc'].unique() if str(c).strip() not in ['', 'nan', 'None']])
        sel_inv_cats = st.multiselect("หมวดหมู่สินค้า (Category)", all_categories, placeholder="ทั้งหมด (All Categories)")

    with col_f2:
        all_statuses = [
            '🔴 สินค้าใกล้หมด (< 14 วัน)',
            '🟢 สต็อกปกติ (14-90 วัน)',
            '🟠 สต็อกล้น (> 90 วัน)',
            '🟡 สินค้าค้างสต็อก (ไม่มียอดขาย 30 วัน)',
            '⚪ หมดสต็อก (0 ชิ้น)'
        ]
        available_statuses = [s for s in all_statuses if s in stock_table['Stock_Status'].values]
        sel_inv_status = st.multiselect("สถานะสุขภาพสต็อก (Health Status)", available_statuses, placeholder="ทั้งหมด (All Statuses)")

    with col_f3:
        search_query = st.text_input("🔍 ค้นหาสินค้า (SKU หรือชื่อสินค้า)", placeholder="พิมพ์รหัส SKU หรือคำในชื่อสินค้า...")

    # Apply Filters
    filtered_stock = stock_table.copy()
    if sel_inv_cats:
        filtered_stock = filtered_stock[filtered_stock['Category_Desc'].isin(sel_inv_cats)]
    if sel_inv_status:
        filtered_stock = filtered_stock[filtered_stock['Stock_Status'].isin(sel_inv_status)]
    if search_query.strip():
        q = search_query.strip().lower()
        filtered_stock = filtered_stock[
            filtered_stock['SKU'].str.lower().str.contains(q, na=False) |
            filtered_stock['Description'].str.lower().str.contains(q, na=False)
        ]

    # ================= Top 5 KPI Scorecards =================
    st.markdown("---")
    kpi1, kpi2, kpi3, kpi4, kpi5 = st.columns(5)
    
    total_qty = filtered_stock['Stock_Available'].sum()
    total_skus = len(filtered_stock)
    total_cost = filtered_stock['Total_Cost_Value'].sum()
    total_expected = filtered_stock['Expected_Sales_Value'].sum()
    total_margin = filtered_stock['Projected_Margin'].sum()
    margin_pct = (total_margin / total_expected * 100) if total_expected > 0 else 0.0

    kpi1.metric("📦 สต็อกคงเหลือรวม", f"{total_qty:,.0f} ชิ้น")
    kpi2.metric("🏷️ จำนวน SKU ในคลัง", f"{total_skus:,.0f} รายการ")
    kpi3.metric("💰 มูลค่าต้นทุนรวม (Cost)", f"฿ {total_cost:,.2f}")
    kpi4.metric("💎 มูลค่าขายคาดหวัง (Sales)", f"฿ {total_expected:,.2f}")
    kpi5.metric("📈 กำไรขั้นต้นคาดการณ์", f"฿ {total_margin:,.2f}", delta=f"{margin_pct:.1f}% Margin")

    st.markdown("---")

    # ================= Visual Analytics (Bar Chart + Donut Chart) =================
    col_c1, col_c2 = st.columns([1.4, 1.0])

    with col_c1:
        # Top 10 Categories by Value
        cat_group = filtered_stock.groupby('Category_Desc').agg({
            'Expected_Sales_Value': 'sum',
            'Total_Cost_Value': 'sum',
            'Stock_Available': 'sum'
        }).reset_index().sort_values('Expected_Sales_Value', ascending=False).head(10)

        fig_cat = go.Figure()
        fig_cat.add_trace(go.Bar(
            x=cat_group['Expected_Sales_Value'],
            y=cat_group['Category_Desc'],
            name='มูลค่าขายคาดหวัง (฿)',
            orientation='h',
            marker=dict(color='#3B82F6'),
            text=[f"฿{v/1e6:.1f}M" if v>=1e6 else f"฿{v/1e3:.0f}k" for v in cat_group['Expected_Sales_Value']],
            textposition='auto'
        ))
        fig_cat.add_trace(go.Bar(
            x=cat_group['Total_Cost_Value'],
            y=cat_group['Category_Desc'],
            name='มูลค่าต้นทุน (฿)',
            orientation='h',
            marker=dict(color='#10B981'),
            text=[f"฿{v/1e6:.1f}M" if v>=1e6 else f"฿{v/1e3:.0f}k" for v in cat_group['Total_Cost_Value']],
            textposition='auto'
        ))
        fig_cat.update_layout(
            barmode='group',
            title='<b>📊 Top 10 หมวดหมู่ตามมูลค่าคงคลัง (มูลค่าขาย vs มูลค่าต้นทุน)</b>',
            yaxis=dict(autorange="reversed"),
            margin=dict(l=10, r=10, t=40, b=10),
            height=360,
            template="plotly_dark",
            paper_bgcolor='rgba(0,0,0,0)',
            plot_bgcolor='rgba(0,0,0,0)',
            legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1)
        )
        st.plotly_chart(fig_cat, use_container_width=True)

    with col_c2:
        # Stock Health Breakdown
        health_counts = filtered_stock['Stock_Status'].value_counts().reset_index()
        health_counts.columns = ['Status', 'Count']
        color_map = {
            '🔴 สินค้าใกล้หมด (< 14 วัน)': '#EF4444',
            '🟢 สต็อกปกติ (14-90 วัน)': '#10B981',
            '🟠 สต็อกล้น (> 90 วัน)': '#F59E0B',
            '🟡 สินค้าค้างสต็อก (ไม่มียอดขาย 30 วัน)': '#EAB308',
            '⚪ หมดสต็อก (0 ชิ้น)': '#9CA3AF'
        }
        colors = [color_map.get(s, '#6B7280') for s in health_counts['Status']]
        fig_health = go.Figure(data=[go.Pie(
            labels=health_counts['Status'],
            values=health_counts['Count'],
            hole=.52,
            marker=dict(colors=colors),
            textinfo='percent+label',
            insidetextorientation='radial'
        )])
        fig_health.update_layout(
            title='<b>🍩 สัดส่วนสุขภาพสต็อก (Stock Health)</b>',
            margin=dict(l=10, r=10, t=40, b=10),
            height=360,
            template="plotly_dark",
            paper_bgcolor='rgba(0,0,0,0)',
            plot_bgcolor='rgba(0,0,0,0)',
            showlegend=False
        )
        st.plotly_chart(fig_health, use_container_width=True)

    # ================= Detailed SKU Grid & Export =================
    col_t_title, col_t_exp = st.columns([3.5, 1.2])
    with col_t_title:
        st.markdown(f"### 📋 รายการสต็อกและมูลค่าสินค้าคงคลังราย SKU ({len(filtered_stock):,d} รายการ)")
    with col_t_exp:
        excel_data = export_inventory_to_excel(filtered_stock)
        st.download_button(
            label="📥 ส่งออก Excel (Stock Report)",
            data=excel_data,
            file_name=f"Stock_Valuation_Report_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="btn_export_stock_excel",
            use_container_width=True
        )

    col_config_inv = {
        "SKU": st.column_config.TextColumn("รหัส SKU", width="medium"),
        "Description": st.column_config.TextColumn("ชื่อสินค้า / รายละเอียด", width="large"),
        "Category_Desc": st.column_config.TextColumn("หมวดหมู่สินค้า", width="medium"),
        "Stock_Status": st.column_config.TextColumn("สถานะสต็อก", width="medium"),
        "Stock_Available": st.column_config.NumberColumn("สต็อกคงเหลือ", format="%,d ชิ้น"),
        "Cost_Unit": st.column_config.NumberColumn("ต้นทุน/ชิ้น", format="฿ %,.2f"),
        "Total_Cost_Value": st.column_config.NumberColumn("มูลค่าต้นทุนรวม", format="฿ %,.2f"),
        "Avg_Selling_Price": st.column_config.NumberColumn("ราคาขายเฉลี่ย", format="฿ %,.2f"),
        "Expected_Sales_Value": st.column_config.NumberColumn("มูลค่าขายคาดหวัง", format="฿ %,.2f"),
        "Units_Sold_30D": st.column_config.NumberColumn("ขายได้ 30 วัน", format="%,d ชิ้น"),
        "Days_of_Inventory": st.column_config.NumberColumn("สต็อกพอขาย", format="%,.0f วัน"),
    }

    display_cols = [
        'SKU', 'Description', 'Category_Desc', 'Stock_Status',
        'Stock_Available', 'Cost_Unit', 'Total_Cost_Value',
        'Avg_Selling_Price', 'Expected_Sales_Value', 'Units_Sold_30D', 'Days_of_Inventory'
    ]

    st.dataframe(
        filtered_stock[display_cols],
        hide_index=True,
        use_container_width=True,
        height=480,
        column_config=col_config_inv
    )
